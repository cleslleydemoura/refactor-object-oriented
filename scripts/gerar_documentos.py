"""Gera diagramas, DOCX/PDF e pacote de entrega a partir de fontes editáveis.

Uso: python scripts/gerar_documentos.py
Para revisar apenas Word/PDF: python scripts/gerar_documentos.py --somente-documentos
Dependências editoriais: scripts/requirements-docs.txt. Não são usadas pelo Java.
"""
from __future__ import annotations
import argparse
import html
import json
import math
import re
import shutil
import sys
import textwrap
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timezone, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".tools"))
sys.path.insert(0, str(ROOT / ".tools/doclibs"))
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image as PdfImage, Preformatted, CondPageBreak

DOCS = ROOT / "docs"
OUT = ROOT / "entrega"
INK = "#203c30"
MUTED = "#627369"
FONTS = Path("C:/Windows/Fonts")

def font_path(filename: str, fallback: str):
    path = FONTS / filename
    if path.exists(): return path
    candidates = [Path("/usr/share/fonts/truetype/dejavu") / fallback,
                  Path("/usr/share/fonts/truetype/liberation2") / filename]
    return next(p for p in candidates if p.exists())

REG = font_path("arial.ttf", "DejaVuSans.ttf")
BOLD = font_path("arialbd.ttf", "DejaVuSans-Bold.ttf")
ITALIC = font_path("ariali.ttf", "DejaVuSans-Oblique.ttf")
MONO = font_path("consola.ttf", "DejaVuSansMono.ttf")

class Diagram:
    """Desenha a mesma figura em SVG vetorial e PNG para documentos Word."""
    def __init__(self, width, height, title):
        self.width, self.height = width, height
        self.img = Image.new("RGB", (width, height), "#fffdf8")
        self.draw = ImageDraw.Draw(self.img)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                    f'<rect width="{width}" height="{height}" fill="#fffdf8"/>']
        self.text(35, 25, title, 30, True)
    def text(self, x, y, value, size=20, bold=False, color=INK):
        font = ImageFont.truetype(str(BOLD if bold else REG), size)
        self.draw.text((x, y), value, font=font, fill=color)
        self.svg.append(f'<text x="{x}" y="{y+size}" font-family="Arial,sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}">{html.escape(value)}</text>')
    def line(self, x1, y1, x2, y2, dashed=False, arrow=False):
        if dashed:
            distance = math.hypot(x2-x1, y2-y1)
            for start in range(0, int(distance), 12):
                end = min(start+7, distance)
                self.draw.line((x1+(x2-x1)*start/distance, y1+(y2-y1)*start/distance,
                                x1+(x2-x1)*end/distance, y1+(y2-y1)*end/distance), fill=MUTED, width=2)
        else:
            self.draw.line((x1, y1, x2, y2), fill=MUTED, width=2)
        dash = 'stroke-dasharray="7 5"' if dashed else ''
        self.svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{MUTED}" stroke-width="2" {dash}/>')
        if arrow:
            angle = math.atan2(y2-y1, x2-x1)
            pts = [(x2, y2), (x2-13*math.cos(angle-.45), y2-13*math.sin(angle-.45)),
                   (x2-13*math.cos(angle+.45), y2-13*math.sin(angle+.45))]
            self.draw.polygon(pts, fill=MUTED)
            points = " ".join(f"{x:.1f},{y:.1f}" for x,y in pts)
            self.svg.append(f'<polygon points="{points}" fill="{MUTED}"/>')
    def rect(self, x, y, w, h, fill="#edf1e2", oval=False):
        if oval: self.draw.ellipse((x,y,x+w,y+h), fill=fill, outline=MUTED, width=2)
        else: self.draw.rounded_rectangle((x,y,x+w,y+h), radius=10, fill=fill, outline=MUTED, width=2)
        if oval: self.svg.append(f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" fill="{fill}" stroke="{MUTED}"/>')
        else: self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="10" fill="{fill}" stroke="{MUTED}"/>')
    def box(self, x, y, w, h, title, lines, fill="#edf1e2"):
        self.rect(x,y,w,h,fill)
        self.text(x+15,y+12,title,22,True)
        self.line(x,y+48,x+w,y+48)
        for i,value in enumerate(lines): self.text(x+15,y+62+i*29,value,18)
    def save(self, name):
        folder = DOCS / "diagramas"
        folder.mkdir(exist_ok=True)
        self.img.save(folder / (name+".png"))
        (folder / (name+".svg")).write_text("\n".join(self.svg+["</svg>"]),encoding="utf-8")

def diagrams():
    d = Diagram(1200,760,"Casos de uso — Acervo")
    d.rect(370,90,770,600,"#f6f4ed")
    d.text(645,108,"Catálogo pessoal",23,True)
    d.draw.ellipse((116,325,164,373), outline=INK, width=3)
    d.svg.append(f'<circle cx="140" cy="349" r="24" fill="none" stroke="{INK}" stroke-width="3"/>')
    for pts in [(140,373,140,460),(95,405,185,405),(140,460,105,510),(140,460,175,510)]: d.line(*pts)
    d.text(45,535,"Usuário do catálogo",20,True)
    labels=["Cadastrar obra","Listar coleção","Buscar título ou autor/diretor","Visualizar detalhes","Editar obra","Excluir com confirmação"]
    for i,label in enumerate(labels):
        y=155+i*82
        d.line(190,405,430,y+31)
        d.rect(430,y,650,62,"#e8eddb",oval=True)
        d.text(455,y+17,label,21)
    d.save("casos_uso")

    d=Diagram(1480,1120,"Classes principais e contratos")
    d.box(35,100,360,170,"ItemMidiaServlet",["- servico: ItemMidiaService","# doGet(request, response)","# doPost(request, response)"])
    d.box(35,350,360,235,"ItemMidiaService",["- dao: ItemMidiaDao","- relogio: Clock","+ salvar(id, campos)","+ buscar(termo) / obter(id)","+ excluir(id)"])
    d.box(530,350,390,235,"«interface» ItemMidiaDao",["+ inserir(item): ItemMidia","+ buscar(termo): List","+ buscarPorId(id): Optional","+ atualizar(item): boolean","+ excluir(id): boolean"])
    d.box(500,710,380,150,"JdbcItemMidiaDao",["- fabrica: FabricaConexao","+ CRUD JDBC parametrizado","- preencher() / mapear()"])
    d.box(960,710,465,150,"ItemMidiaDaoEmMemoria",["Somente em src/test","Implementa o mesmo contrato","Substitui JDBC nos testes de serviço"])
    d.box(980,100,440,315,"ItemMidia",["- id: Long","- titulo, autorDiretor: String","- anoLancamento: int","- genero, sinopse: String","- tipoMidia: TipoMidia","+ getters / comId(id)","Estado privado e imutável"])
    d.box(980,460,440,130,"«enum» TipoMidia",["LIVRO / FILME / SERIE","+ getDescricao(): String"])
    d.box(35,945,400,130,"«interface» FabricaConexao",["+ abrir(): Connection","Contrato de criação de conexões"])
    d.box(530,945,390,130,"ConexaoJdbc",["URL, usuário e senha","DriverManager.getConnection"])
    d.line(215,270,215,350,arrow=True)
    d.line(395,465,530,465,arrow=True)
    d.line(690,710,690,585,dashed=True,arrow=True)
    d.line(1190,710,870,585,dashed=True,arrow=True)
    d.line(980,415,980,460,arrow=True)
    d.line(395,375,980,240,dashed=True,arrow=True)
    d.line(500,800,350,945,arrow=True)
    d.line(530,1010,435,1010,dashed=True,arrow=True)
    d.text(40,630,"Seta contínua: dependência. Tracejada entre implementação e interface: realização.",19)
    d.text(40,660,"A ligação do serviço com a entidade representa uso do tipo de domínio.",19)
    d.save("classes")

    d=Diagram(1100,690,"Modelo de dados — uma entidade principal")
    d.box(100,105,900,460,"ITEM_MIDIA",[
        "PK  id                  BIGINT IDENTITY",
        "     titulo              VARCHAR(255) NOT NULL",
        "     autor_diretor        VARCHAR(255) NOT NULL",
        "     ano_lancamento       INTEGER NOT NULL; CHECK 1..9999",
        "     genero              VARCHAR(100) NOT NULL",
        "     sinopse             VARCHAR(2000) NOT NULL",
        "     tipo_midia          VARCHAR(10) NOT NULL",
        "CHECK tipo_midia IN ('LIVRO', 'FILME', 'SERIE')",
        "CHECK CHAR_LENGTH(TRIM(texto)) > 0 nos campos textuais"])
    d.text(100,600,"Não há chaves estrangeiras: o escopo atual utiliza uma única tabela.",20)
    d.save("der")

    d=Diagram(1500,1000,"Sequência — cadastro válido")
    names=["Navegador","Filtro CSRF","Servlet","Serviço","DAO JDBC","H2"]
    xs=[125,375,625,875,1125,1375]
    for x,name in zip(xs,names):
        d.rect(x-100,90,200,65)
        d.text(x-88,110,name,22,True)
        d.line(x,155,x,920,dashed=True)
    msgs=[(0,1,"POST salvar + campos + token"),(1,2,"Token válido"),(2,3,"salvar(null, campos)"),
          (3,3,"Normaliza e valida"),(3,4,"inserir(item)"),(4,5,"INSERT com parâmetros"),
          (5,4,"Chave gerada"),(4,3,"Item com ID"),(3,2,"Item persistido"),(2,0,"303 → GET detalhes"),
          (0,2,"GET detalhes?id=..."),(2,0,"HTML renderizado pela JSP")]
    for i,(a,b,label) in enumerate(msgs):
        y=205+i*59
        if a==b:
            d.line(xs[a],y,xs[a]+140,y);d.line(xs[a]+140,y,xs[a]+140,y+22);d.line(xs[a]+140,y+22,xs[a],y+22,arrow=True)
            d.text(xs[a]+10,y-28,label,16)
        else:
            d.line(xs[a],y,xs[b],y,arrow=True)
            d.text(min(xs[a],xs[b])+10,y-26,label,17)
    d.text(35,947,"O GET de detalhes consulta serviço/DAO/banco antes de renderizar a JSP (fluxo completo no fonte Mermaid).",18)
    d.save("sequencia")

def verification():
    suites=[]
    folders=[ROOT/"target/surefire-reports",ROOT/"target/failsafe-reports"]
    for folder in folders:
        for path in sorted(folder.glob("TEST-*.xml")):
            suites.append(ET.parse(path).getroot())
    expected={"br.edu.catalogo.dao.JdbcItemMidiaDaoTest","br.edu.catalogo.service.ItemMidiaServiceTest","br.edu.catalogo.web.CatalogoHttpIT"}
    if {x.get("name") for x in suites} != expected: raise RuntimeError("Execute maven verify antes de gerar a entrega: faltam suítes.")
    totals={k:sum(int(s.get(k,"0")) for s in suites) for k in ["tests","failures","errors","skipped"]}
    if totals["failures"] or totals["errors"] or totals["skipped"]: raise RuntimeError("A entrega exige todas as suítes aprovadas sem testes ignorados.")
    lines=["# Registro de verificação executada", "", "Data da preparação: 17 de outubro de 2026 (America/Sao_Paulo).", "",
           "Comando: `maven.ps1 verify` (Maven 3.9.9; JDK 17.0.12).", "",
           "Resultado: **APROVADO**. As contagens abaixo correspondem aos relatórios XML JUnit da execução registrada na entrega.", "",
           "| Suíte | Testes | Falhas | Erros | Ignorados | Duração da suíte |", "| --- | --- | --- | --- | --- | --- |"]
    for s in suites:
        lines.append(f'| {s.get("name").rsplit(".",1)[-1]} | {s.get("tests")} | {s.get("failures")} | {s.get("errors")} | {s.get("skipped")} | {s.get("time")} s |')
    lines += [f'| Total | {totals["tests"]} | {totals["failures"]} | {totals["errors"]} | {totals["skipped"]} | — |',"",
              "Os tempos da tabela correspondem à execução das suítes de testes. O desenvolvimento do projeto foi realizado ao longo de dez dias.","",
              "## Ambiente e alcance", "", "Windows; Java 17; Tomcat 10.1.61; H2 2.3.232; Servlet 6.0; JSP 3.1; JUnit 5.12.2. Testes JDBC usam bancos isolados. Testes HTTP iniciam um Tomcat real em porta livre e exercitam as JSPs.", "",
              "Foi verificada persistência real com fechamento e reabertura de um banco H2 em arquivo temporário. A indisponibilidade simulada do banco produz 503 e mensagem pública sem detalhes internos. O log técnico dessa simulação é esperado e não representa uma falha da suíte.", "",
              "## Casos executados", ""]
    for s in suites:
        lines += ["### "+s.get("name").rsplit(".",1)[-1],""]
        sanitized=ET.Element("testsuite", {k:s.get(k,"") for k in ["name","tests","failures","errors","skipped","time"]})
        for case in s.findall("testcase"):
            lines.append("- APROVADO: "+case.get("name"))
            ET.SubElement(sanitized,"testcase", {k:case.get(k,"") for k in ["name","classname","time"]})
        lines.append("")
        dest=OUT/"evidencias_testes"/(s.get("name").rsplit(".",1)[-1]+".xml")
        dest.parent.mkdir(parents=True,exist_ok=True)
        ET.ElementTree(sanitized).write(dest,encoding="utf-8",xml_declaration=True)
    lines += ["## Escopo da avaliação", "", "A avaliação abrange as regras de negócio, o acesso JDBC, a persistência em arquivo e os fluxos HTTP com Tomcat embutido. Os testes verificam também entradas inválidas, CSRF, escape de HTML e respostas de erro. Testes de carga, avaliação com usuários e análise formal de acessibilidade são possibilidades de aprofundamento do projeto."]
    (DOCS/"VERIFICACAO.md").write_text("\n".join(lines),encoding="utf-8")
    return totals

def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1;continue
        if line.startswith("```"):
            language=line[3:];i+=1;body=[]
            while i<len(lines) and not lines[i].startswith("```"): body.append(lines[i]);i+=1
            yield ("code", "\n".join(body));i+=1;continue
        if line.startswith("#"):
            m=re.match(r"(#+)\s+(.*)",line);yield ("heading",(len(m.group(1)),m.group(2)));i+=1;continue
        image=re.fullmatch(r"!\[(.*?)\]\((.*?)\)",line)
        if image: yield ("image",(image.group(1),DOCS/image.group(2)));i+=1;continue
        if line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                cells=[c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",c.replace(" ","")) for c in cells): rows.append(cells)
                i+=1
            yield ("table",rows);continue
        if re.match(r"^(- |\d+\. )",line):
            ordered=bool(re.match(r"\d+\.",line));yield ("list",(ordered,re.sub(r"^(- |\d+\. )","",line)));i+=1;continue
        body=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#|```|\||!\[|- |\d+\. )",lines[i]): body.append(lines[i]);i+=1
        yield ("paragraph"," ".join(body))

def plain(value):
    value=re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",value)
    value=value.replace("**", "").replace("`", "")
    return value

def pdf_markup(value):
    value=value.replace("☐","[ ]").replace("☒","[X]")
    value=html.escape(value)
    value=re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)",r'<link href="\2" color="#315c46">\1</link>',value)
    value=re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",value)
    value=re.sub(r"\*\*(.+?)\*\*",r"<b>\1</b>",value)
    value=re.sub(r"`([^`]+)`",r'<font name="Mono">\1</font>',value)
    return value

def add_word_runs(p,text):
    parts=re.split(r"(\*\*.*?\*\*|`[^`]+`)",text)
    for part in parts:
        run=p.add_run(plain(part))
        if part.startswith("**"):run.bold=True
        if part.startswith("`"):run.font.name="Consolas";run.font.size=Pt(9)

def write_docx(path, source, cover=False):
    doc=Document();sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.6);sec.right_margin=Cm(2.2)
    normal=doc.styles["Normal"];normal.font.name="Arial";normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.2
    for name in ["Title","Heading 1","Heading 2","Heading 3"]:
        doc.styles[name].font.name="Arial";doc.styles[name].font.color.rgb=RGBColor.from_string("203C30")
    sec.header.paragraphs[0].text="ACERVO  |  Situação-Problema 3"
    sec.header.paragraphs[0].style=doc.styles["Caption"]
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run("Página ")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    for kind,value in blocks(source):
        if kind=="heading":
            level,title=value
            if level==1 and cover and len(doc.paragraphs)>1:doc.add_page_break()
            doc.add_heading(plain(title),min(level,3))
        elif kind=="paragraph":add_word_runs(doc.add_paragraph(),value)
        elif kind=="list":
            ordered,text=value;add_word_runs(doc.add_paragraph(style="List Number" if ordered else "List Bullet"),text)
        elif kind=="code":
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10)
            r=p.add_run(value);r.font.name="Consolas";r.font.size=Pt(8)
        elif kind=="table":
            table=doc.add_table(rows=0,cols=len(value[0]));table.style="Light Shading Accent 1"
            for ri,row in enumerate(value):
                cells=table.add_row().cells
                for ci,text in enumerate(row):
                    cells[ci].text=plain(text)
                    for p in cells[ci].paragraphs:
                        for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
            doc.add_paragraph()
        elif kind=="image":
            caption,image=value
            if not image.exists():raise FileNotFoundError(image)
            doc.add_picture(str(image),width=Cm(15.9))
            doc.add_paragraph(caption,style="Caption")
    doc.core_properties.title=path.stem
    doc.core_properties.subject="Refatoração e Orientação a Objetos"
    doc.core_properties.author=""
    doc.save(path)

def register_fonts():
    for name,path in [("Body",REG),("BodyBold",BOLD),("BodyItalic",ITALIC),("Mono",MONO)]:pdfmetrics.registerFont(TTFont(name,str(path)))
    pdfmetrics.registerFontFamily("Body",normal="Body",bold="BodyBold",italic="BodyItalic",boldItalic="BodyBold")

def write_pdf(path,source,cover=False):
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Text",fontName="Body",fontSize=10.2,leading=15.2,spaceAfter=9,alignment=TA_JUSTIFY,splitLongWords=True))
    for n,size,leading in [("H1",21,27),("H2",15,20),("H3",12,17)]:
        styles.add(ParagraphStyle(name=n,fontName="BodyBold",fontSize=size,leading=leading,spaceBefore=13,spaceAfter=10,textColor=colors.HexColor(INK),keepWithNext=True))
    styles.add(ParagraphStyle(name="Cell",fontName="Body",fontSize=8,leading=11,spaceAfter=0,splitLongWords=True))
    styles.add(ParagraphStyle(name="CaptionSmall",fontName="BodyItalic",fontSize=8.2,leading=11,spaceAfter=12,textColor=colors.HexColor(MUTED)))
    story=[]
    for kind,value in blocks(source):
        if kind=="heading":
            level,title=value
            if level==1 and cover and story:story.append(PageBreak())
            story.append(Paragraph(pdf_markup(title),styles["H"+str(min(level,3))]))
        elif kind=="paragraph":story.append(Paragraph(pdf_markup(value),styles["Text"]))
        elif kind=="list":
            ordered,text=value
            story.append(Paragraph(pdf_markup(text),styles["Text"],bulletText="•"))
        elif kind=="code":
            wrapped="\n".join("\n".join(textwrap.wrap(line,99,replace_whitespace=False,drop_whitespace=False) or [""]) for line in value.splitlines())
            code_style=ParagraphStyle("Code",fontName="Mono",fontSize=7.5,leading=10,spaceBefore=6,spaceAfter=12,backColor=colors.HexColor("#f0f2e9"),borderPadding=7)
            story.append(Preformatted(wrapped,code_style))
        elif kind=="table":
            data=[[Paragraph(pdf_markup(c),styles["Cell"]) for c in row] for row in value]
            widths=[(A4[0]-135)/len(value[0])]*len(value[0])
            table=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e3ead5")),("VALIGN",(0,0),(-1,-1),"TOP"),
                ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#cbd4c3")),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
            story+=[table,Spacer(1,12)]
        elif kind=="image":
            caption,image=value
            if not image.exists():raise FileNotFoundError(image)
            with Image.open(image) as im:w,h=im.size
            maxw=A4[0]-135;maxh=520
            factor=min(maxw/w,maxh/h)
            story += [CondPageBreak(min(h*factor+45,560)),PdfImage(str(image),width=w*factor,height=h*factor),Spacer(1,7),Paragraph(html.escape(caption),styles["CaptionSmall"])]
    def decorate(canvas,document):
        canvas.saveState();canvas.setStrokeColor(colors.HexColor("#d6dece"));canvas.line(68,A4[1]-43,A4[0]-67,A4[1]-43)
        canvas.setFont("Body",8);canvas.setFillColor(colors.HexColor(MUTED));canvas.drawString(68,A4[1]-34,"ACERVO  |  Refatoração e Orientação a Objetos")
        canvas.drawRightString(A4[0]-67,32,str(document.page));canvas.restoreState()
    document=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=67,leftMargin=68,topMargin=64,bottomMargin=54,title=path.stem,author="",pageCompression=1)
    document.build(story,onFirstPage=decorate,onLaterPages=decorate)

def sources_pdf():
    parts=["# Código-fonte de referência", "", "Versão final da implementação. As quebras de linhas longas são apenas editoriais; os arquivos originais em src são a referência executável. Cada linha original começa com seu número. O guia do código explica as responsabilidades e os métodos."]
    paths=sorted((ROOT/"src").rglob("*.java"))+[ROOT/"pom.xml",ROOT/"src/main/resources/db/schema.sql",ROOT/"src/main/webapp/WEB-INF/web.xml"]+sorted((ROOT/"src/main/webapp/WEB-INF/views").glob("*.jsp*"))+[ROOT/"src/main/webapp/assets/app.js",ROOT/"maven.ps1",ROOT/"iniciar.ps1"]
    for path in paths:
        parts += ["", "## "+path.relative_to(ROOT).as_posix(),"", "```text", *[f"{i:03} | {line}" for i,line in enumerate(path.read_text(encoding="utf-8").splitlines(),1)],"```"]
    write_pdf(OUT/"Codigo_Fonte_Referencia.pdf","\n".join(parts))

def package():
    shutil.copy2(ROOT/"target/catalogo.war",OUT/"catalogo.war")
    tracked=[]
    for folder in ["src","docs","scripts"]: tracked.extend(p for p in (ROOT/folder).rglob("*") if p.is_file() and "__pycache__" not in p.parts)
    tracked += [ROOT/p for p in ["README.md","pom.xml","maven.ps1","iniciar.ps1",".gitignore"]]
    documents=[p for p in OUT.rglob("*") if p.is_file() and p.suffix!=".zip"]
    archive=OUT/"Acervo_Situacao_Problema_3_Completo.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for path in tracked+documents:z.write(path,"acervo/"+path.relative_to(ROOT).as_posix())
        for folder in ["target/classes","target/dependency","target/catalogo"]:
            for path in (ROOT/folder).rglob("*"):
                if path.is_file():z.write(path,"acervo/"+path.relative_to(ROOT).as_posix())
    return archive

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--somente-documentos", action="store_true",
                        help="Atualiza Word/PDF a partir dos Markdown, preservando os registros de testes e a aplicação.")
    args = parser.parse_args()
    OUT.mkdir(exist_ok=True);register_fonts()
    if not args.somente_documentos:
        diagrams()
        totals=verification()
    names=[("PROJETO_INTERVENCAO","Projeto_Intervencao_Preenchido"),("RELATORIO_TECNICO","Relatorio_Tecnico"),
           ("GUIA_CODIGO","Guia_Detalhado_Codigo"),("MANUAL_USUARIO","Manual_Usuario"),("ROTEIRO_APRESENTACAO","Roteiro_Apresentacao"),("VERIFICACAO","Verificacao_Testes")]
    combined=[]
    for source,output in names:
        content=(DOCS/(source+".md")).read_text(encoding="utf-8")
        write_docx(OUT/(output+".docx"),content);write_pdf(OUT/(output+".pdf"),content)
        combined.append(content)
    full="\n\n".join(combined)
    (DOCS/"DOCUMENTACAO_COMPLETA.md").write_text(full,encoding="utf-8")
    write_docx(OUT/"Documentacao_Completa.docx",full,cover=True)
    write_pdf(OUT/"Documentacao_Completa.pdf",full,cover=True)
    if not args.somente_documentos:
        sources_pdf();archive=package()
    from pypdf import PdfReader
    summary={p.name:{"paginas":len(PdfReader(str(p)).pages),"bytes":p.stat().st_size} for p in sorted(OUT.glob("*.pdf"))}
    result={"documentos":summary}
    if not args.somente_documentos:
        result.update({"testes":totals,"pacote":str(archive),"tamanho_zip":archive.stat().st_size})
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=="__main__":main()
